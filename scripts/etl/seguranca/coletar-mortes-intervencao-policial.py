# -*- coding: utf-8 -*-
"""
Coletor das MORTES POR INTERVENÇÃO DE AGENTE DO ESTADO (letalidade policial) —
Sinesp VDE / MJSP.

Papel no portal
===============
Alimenta `/direitos-em-movimento/seguranca-publica`: quantas pessoas morreram
por intervenção de agente do Estado no Brasil, por UF e por sexo, e que fatia
isso representa do conjunto das mortes violentas intencionais. É o dado que
torna visível a letalidade policial — recorde histórico no país.

Fonte oficial
=============
MJSP — Dados Nacionais de Segurança Pública (Sinesp VDE), base `bancovde-2025.xlsx`:
https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/dados-nacionais-1

Cortesia e download
===================
O gov.br responde 403 a clientes HTTP de biblioteca (WAF): download via `curl`,
com cache em disco. User-Agent honesto; um arquivo.

Números
=======
A base traz, por UF/mês, o evento "Morte por intervenção de Agente do Estado"
com a contagem por sexo. O total geral e a participação nas mortes violentas
intencionais (homicídio doloso + latrocínio + lesão corporal seguida de morte +
morte por intervenção) são calculados aqui — nunca digitados à mão.
"""

import json
import re
import subprocess
import sys
import time
from collections import defaultdict
from pathlib import Path

import openpyxl

URL_FONTE = (
    "https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/"
    "dados-nacionais-1/base-de-dados-e-notas-metodologicas-dos-gestores-estaduais-sinesp-vde-2022-e-2023"
)
URL_ARQUIVO = (
    "https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/"
    "download/dnsp-base-de-dados/bancovde-2025.xlsx/@@download/file"
)
USER_AGENT = "ControlePopular/1.0 (+contato@controlepopular.com.br; dado publico governamental)"
ENTRADA = Path("etl/betim/entrada/sinesp/bancovde-2025.xlsx")
SAIDA = Path("apps/web/data/seguranca/mortes-intervencao-policial-2025.json")

EVENTO_MDIP = "Morte por intervenção de Agente do Estado"
# Definição de "mortes violentas intencionais" usada pelo FBSP.
EVENTOS_MVI = ["Homicídio doloso", "Roubo seguido de morte (latrocínio)", "Lesão corporal seguida de morte"]


def baixar(destino: Path) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.exists() and destino.stat().st_size > 5_000_000:
        print(f"[sinesp] XLSX em disco ({destino.stat().st_size} bytes) — checkpoint usado")
        return destino
    print("[sinesp] baixando (via curl) …")
    subprocess.run(
        ["curl.exe", "-s", "-L", "-o", str(destino), "-A", USER_AGENT, "-e", URL_FONTE, URL_ARQUIVO],
        check=True,
    )
    print(f"[sinesp] ok ({destino.stat().st_size} bytes)")
    time.sleep(2)
    return destino


def sem_acento(s: str) -> str:
    import unicodedata

    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def n(v) -> int:
    try:
        return int(str(v or "0").strip() or 0)
    except ValueError:
        return 0


def main() -> int:
    caminho = baixar(ENTRADA)
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    por_uf: dict[str, dict] = defaultdict(lambda: {"mdip": 0, "fem": 0, "masc": 0, "ni": 0, "mvi": 0, "meses": {}})
    mdip_nac = {"fem": 0, "masc": 0, "ni": 0}
    mvi_nac = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 12:
            continue
        uf = str(row[0] or "").strip()
        evento = sem_acento(str(row[2] or "")).lower()
        if len(uf) != 2:
            continue
        fem, masc, ni = n(row[7]), n(row[8]), n(row[9])
        if "intervencao de agente do estado" in evento:
            a = por_uf[uf]
            a["mdip"] += fem + masc + ni
            a["fem"] += fem
            a["masc"] += masc
            a["ni"] += ni
            mdip_nac["fem"] += fem
            mdip_nac["masc"] += masc
            mdip_nac["ni"] += ni
            mes = str(row[3])[:7]
            a["meses"][mes] = a["meses"].get(mes, 0) + fem + masc + ni
        elif any(sem_acento(ev).lower() in evento for ev in EVENTOS_MVI):
            por_uf[uf]["mvi"] += fem + masc + ni
            mvi_nac += fem + masc + ni

    mvi_nac += sum(mdip_nac.values())
    registros = []
    for uf, a in por_uf.items():
        mvi = a["mvi"] + a["mdip"]
        registros.append(
            {
                "uf": uf,
                "mdip": a["mdip"],
                "feminino": a["fem"],
                "masculino": a["masc"],
                "nao_informado": a["ni"],
                "mvi": mvi,
                "participacao_pct": round(100 * a["mdip"] / mvi, 1) if mvi else None,
                "serie_mensal": dict(sorted(a["meses"].items())),
            }
        )
    registros.sort(key=lambda r: -r["mdip"])
    mdip_total = sum(mdip_nac.values())

    acervo = {
        "fonte": "MJSP — Dados Nacionais de Segurança Pública (Sinesp VDE), base bancovde-2025",
        "url_fonte": URL_FONTE,
        "url_arquivo": URL_ARQUIVO,
        "ano": 2025,
        "atualizado_em": time.strftime("%Y-%m-%d"),
        "metodologia": (
            "Contagem das mortes por intervenção de agente do Estado reportadas ao Sinesp VDE, por UF e mês. "
            "A participação é calculada sobre as mortes violentas intencionais (homicídio doloso + latrocínio + "
            "lesão corporal seguida de morte + morte por intervenção de agente do Estado), definição usada pelo "
            "FBSP. O dado é agregado — o Sinesp não publica nome, CPF nem identificação de vítima ou agente. "
            "A base não traz recorte racial; o FBSP publica esse recorte no Anuário."
        ),
        "brasil": {
            "mdip": mdip_total,
            "feminino": mdip_nac["fem"],
            "masculino": mdip_nac["masc"],
            "nao_informado": mdip_nac["ni"],
            "mvi": mvi_nac,
            "participacao_pct": round(100 * mdip_total / mvi_nac, 1) if mvi_nac else None,
        },
        "total_ufs": len(registros),
        "registros": registros,
    }
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(json.dumps(acervo, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"ok: MDIP {mdip_total:,} ({acervo['brasil']['participacao_pct']}% das MVI {mvi_nac:,}) -> {SAIDA}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
