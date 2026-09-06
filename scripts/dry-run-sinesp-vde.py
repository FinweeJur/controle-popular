"""Teste dry-run do ETL SINESP VDE — lê xlsx, mapeia IBGE7, não grava no banco."""
import csv
import os
import sys
import unicodedata
from collections import Counter
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DADOS = os.path.join(RAIZ, "etl", "betim", "dados", "temp")

def sem_acento(s):
    nfd = unicodedata.normalize("NFD", s)
    return "".join(c for c in nfd if unicodedata.category(c) != "Mn")

def normalizar(nome):
    return sem_acento(nome).strip().lower()

# Mapeamento
mapa = {}
with open(os.path.join(RAIZ, "apps", "web", "data", "mapeamento_municipios.csv"), encoding="utf-8") as f:
    for r in csv.reader(f, delimiter="|"):
        if r[0] != "municipio_sem_acento":
            mapa[r[0]] = r[1]
print(f"Mapeamento: {len(mapa)} municipios")

# Ler xlsx
arquivos = sorted(f for f in os.listdir(DADOS) if f.startswith("BancoVDE") and f.endswith(".xlsx"))
for arq in arquivos:
    path = os.path.join(DADOS, arq)
    print(f"\n--- {arq} ---")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    total = 0
    mapeados = 0
    sem_map = Counter()
    eventos = Counter()
    meses = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 11 or row[0] != "MG":
            continue
        total += 1
        mun = (row[1] or "").strip()
        evento = (row[2] or "").strip()
        eventos[evento] += 1
        if row[3]:
            meses[row[3].month] += 1
        norm = normalizar(mun)
        if norm in mapa:
            mapeados += 1
        else:
            sem_map[mun] += 1
    wb.close()
    print(f"  Linhas MG: {total}")
    print(f"  Mapeadas: {mapeados} ({100*mapeados//max(1,total)}%)")
    print(f"  Sem mapeamento: {total-mapeados}")
    if sem_map:
        print(f"  Nomes sem match: {dict(sem_map.most_common(5))}")
    print(f"  Eventos: {dict(eventos.most_common())}")
    print(f"  Meses: {sorted(meses.keys())}")
